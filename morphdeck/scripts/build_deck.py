#!/usr/bin/env python3
"""Build a cinematic PowerPoint deck from a morphdeck JSON spec.

Every slide shares a "stage" of decorative shapes named with the `!!` prefix.
PowerPoint's Morph transition matches shapes by that name and glides them
between their per-layout poses, so the whole deck plays like one continuous
camera move. Content on each slide then builds in with timed entrance
animations (rise, fade, zoom, wipe, wheel, letter-by-letter) that start
automatically after the transition.

Usage:
    python3 build_deck.py spec.json out.pptx [--theme NAME] [--motion calm|normal|dramatic]
                                             [--aspect 16:9|16:10|4:3]
    python3 build_deck.py --list-themes
"""
import argparse
import itertools
import json
import math
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import images  # noqa: E402

# Layouts are designed on a 16:9 canvas of DW x DH inches; other aspect ratios
# are mapped onto it (round shapes stay round).
DW, DH = 13.333, 7.5
ASPECTS = {"16:9": (13.333, 7.5), "16:10": (12.0, 7.5), "4:3": (10.0, 7.5)}

NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_DECL = f'xmlns:p="{NS_P}" xmlns:a="{NS_A}"'

# Motion profiles: scale timing, travel distance and ambient movement.
MOTION = {
    "calm":     dict(tempo=1.3, stagger=1.35, dur=1.25, dy=0.035, zoom=0.8, letters=False,
                     spin=35, sway=5, float=0.012, kenburns=104000),
    "normal":   dict(tempo=1.0, stagger=1.0, dur=1.0, dy=0.06, zoom=0.55, letters=True,
                     spin=70, sway=12, float=0.025, kenburns=108000),
    "dramatic": dict(tempo=0.85, stagger=0.8, dur=0.9, dy=0.13, zoom=0.25, letters=True,
                     spin=140, sway=25, float=0.045, kenburns=115000),
}

# Stage poses per layout, in design inches: (x, y, w, h).
# ga/gb = glows, ring = dashed ring, disc = small solid dot, panel = surface/accent
# block, bar = accent bar. Morph tweens between these on every slide change.
POSES = {
    "title":     dict(ga=(6.6, -1.8, 8.6, 8.6), gb=(-2.6, 3.6, 5.6, 5.6), ring=(8.4, 0.7, 5.4, 5.4),
                      disc=(11.6, 5.7, 0.6, 0.6), panel=(0, 0, 0.16, DH), bar=(0.9, 5.05, 1.5, 0.09)),
    "section":   dict(ga=(-2.5, -2.0, 9.5, 9.5), gb=(9.5, 4.0, 5.0, 5.0), ring=(0.9, 1.25, 5.0, 5.0),
                      disc=(5.4, 1.35, 0.5, 0.5), panel=(0, 0, 0.16, DH), bar=(6.6, 4.25, 1.2, 0.09)),
    "agenda":    dict(ga=(-3.0, 2.5, 8.0, 8.0), gb=(10.0, -3.0, 5.0, 5.0), ring=(-1.4, 3.6, 4.6, 4.6),
                      disc=(5.15, 0.95, 0.45, 0.45), panel=(5.85, 0.9, 0.04, 5.7), bar=(0.9, 3.35, 1.0, 0.09)),
    "statement": dict(ga=(-3.2, -3.2, 9.0, 9.0), gb=(9.2, 3.4, 6.0, 6.0), ring=(-1.9, 4.2, 4.6, 4.6),
                      disc=(12.0, 0.9, 0.45, 0.45), panel=(1.3, 1.15, 10.7, 5.2), bar=(1.95, 1.75, 0.9, 0.09)),
    "question":  dict(ga=(3.2, -0.8, 7.0, 7.0), gb=(-2.8, -2.8, 5.0, 5.0), ring=(3.92, 0.5, 5.5, 5.5),
                      disc=(10.4, 1.0, 0.4, 0.4), panel=(0, 0, 0.16, DH), bar=(6.17, 5.25, 1.0, 0.09)),
    "quote":     dict(ga=(7.6, 0.0, 7.5, 7.5), gb=(-3.0, 4.6, 5.0, 5.0), ring=(8.65, 1.15, 4.3, 4.3),
                      disc=(12.25, 1.2, 0.5, 0.5), panel=(0, 0, 0.16, DH), bar=(0.9, 5.25, 0.9, 0.09)),
    "bullets":   dict(ga=(7.6, 0.6, 8.0, 8.0), gb=(-3.0, -3.0, 5.5, 5.5), ring=(9.1, 1.45, 4.6, 4.6),
                      disc=(9.0, 1.2, 0.55, 0.55), panel=(8.9, 0, DW - 8.9, DH), bar=(0.9, 2.0, 1.0, 0.09)),
    "stat":      dict(ga=(6.8, -0.6, 8.0, 8.0), gb=(-3.0, 4.2, 5.5, 5.5), ring=(7.9, 0.95, 5.6, 5.6),
                      disc=(12.1, 0.8, 0.5, 0.5), panel=(0, 0, 0.16, DH), bar=(0.9, 5.05, 1.2, 0.09)),
    "chart":     dict(ga=(8.2, 1.2, 7.0, 7.0), gb=(-3.0, -3.5, 5.5, 5.5), ring=(10.1, -1.2, 3.2, 3.2),
                      disc=(12.4, 6.3, 0.4, 0.4), panel=(9.45, 2.0, 3.0, 4.75), bar=(0.9, 1.75, 1.0, 0.09)),
    "cards":     dict(ga=(8.5, -4.0, 7.5, 7.5), gb=(-2.5, 4.4, 5.5, 5.5), ring=(10.9, -1.6, 3.4, 3.4),
                      disc=(12.25, 0.55, 0.4, 0.4), panel=(0, 6.95, DW, 0.55), bar=(0.9, 1.85, 1.0, 0.09)),
    "people":    dict(ga=(-2.0, -4.2, 7.0, 7.0), gb=(9.0, 4.2, 6.0, 6.0), ring=(11.6, 5.4, 3.0, 3.0),
                      disc=(0.45, 0.5, 0.35, 0.35), panel=(0, 0, DW, 0.1), bar=(0.9, 1.85, 1.0, 0.09)),
    "timeline":  dict(ga=(-3.5, -4.5, 8.0, 8.0), gb=(9.6, 3.9, 5.8, 5.8), ring=(11.2, -1.4, 3.4, 3.4),
                      disc=(0.62, 4.13, 0.3, 0.3), panel=(0, 0, 0.16, DH), bar=(0.9, 4.24, 11.6, 0.05)),
    "process":   dict(ga=(7.2, -0.2, 7.0, 7.0), gb=(-3.0, 4.3, 5.0, 5.0), ring=(7.65, 1.05, 5.2, 5.2),
                      disc=(12.55, 0.6, 0.4, 0.4), panel=(0, 0, 0.16, DH), bar=(0.9, 1.95, 1.0, 0.09)),
    "compare":   dict(ga=(3.6, -4.5, 6.0, 6.0), gb=(3.6, 5.0, 6.0, 6.0), ring=(5.97, 3.37, 1.4, 1.4),
                      disc=(6.47, 3.87, 0.4, 0.4), panel=(6.635, 2.05, 0.06, 4.7), bar=(0.9, 1.85, 1.0, 0.09)),
    "split":     dict(ga=(-3.0, 3.5, 6.5, 6.5), gb=(8.0, -1.5, 7.5, 7.5), ring=(7.6, 1.5, 4.5, 4.5),
                      disc=(7.25, 5.9, 0.5, 0.5), panel=(7.0, 0, DW - 7.0, DH), bar=(0.9, 3.3, 1.0, 0.09)),
    "image":     dict(ga=(4.0, -1.0, 9.0, 9.0), gb=(-2.0, 3.0, 6.0, 6.0), ring=(1.0, 0.5, 6.5, 6.5),
                      disc=(12.0, 6.5, 0.4, 0.4), panel=(0, 0, DW, DH), bar=(0.9, 4.6, 1.0, 0.09)),
    "gallery":   dict(ga=(9.0, -4.0, 7.0, 7.0), gb=(-3.0, 4.5, 5.5, 5.5), ring=(-1.2, -1.6, 3.2, 3.2),
                      disc=(12.4, 0.65, 0.35, 0.35), panel=(0, 6.95, DW, 0.55), bar=(0.9, 1.75, 1.0, 0.09)),
    "closing":   dict(ga=(2.6, -2.0, 8.2, 8.2), gb=(-2.5, 4.5, 5.0, 5.0), ring=(3.67, 0.25, 6.0, 6.0),
                      disc=(9.3, 1.0, 0.55, 0.55), panel=(0, 0, 0.16, DH), bar=(6.17, 4.4, 1.0, 0.09)),
}
LAYOUTS = tuple(POSES)
ACCENT_PANEL = ("compare", "title", "section", "stat", "timeline", "closing", "agenda",
                "question", "quote", "process")


def emu(inches):
    return Emu(int(round(inches * 914400)))


def load_themes():
    with open(os.path.join(HERE, "themes.json")) as f:
        return json.load(f)


# ---------------------------------------------------------------------------
# Low-level XML helpers
# ---------------------------------------------------------------------------

FILL_TAGS = ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill")


def _strip_style(shape):
    """Drop the theme style python-pptx attaches (it adds a drop shadow)."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def _set_fill_xml(shape, fill_xml):
    sppr = shape._element.spPr
    for tag in FILL_TAGS:
        for el in sppr.findall(qn(tag)):
            sppr.remove(el)
    fill = etree.fromstring(f'<root xmlns:a="{NS_A}">{fill_xml}</root>')[0]
    geom = sppr.find(qn("a:prstGeom"))
    if geom is None:
        geom = sppr.find(qn("a:xfrm"))
    geom.addnext(fill)


def _clr(hex_, alpha=1.0):
    if alpha >= 1:
        return f'<a:srgbClr val="{hex_}"/>'
    return f'<a:srgbClr val="{hex_}"><a:alpha val="{int(alpha * 100000)}"/></a:srgbClr>'


def solid(shape, hex_, alpha=1.0):
    _set_fill_xml(shape, f"<a:solidFill>{_clr(hex_, alpha)}</a:solidFill>")


def glow(shape, hex_, alpha):
    """Radial gradient from `alpha` at the centre to transparent at the edge."""
    _set_fill_xml(shape, (
        '<a:gradFill rotWithShape="1"><a:gsLst>'
        f'<a:gs pos="0">{_clr(hex_, alpha)}</a:gs>'
        f'<a:gs pos="45000">{_clr(hex_, alpha * 0.45)}</a:gs>'
        f'<a:gs pos="100000">{_clr(hex_, 0.0001)}</a:gs>'
        '</a:gsLst><a:path path="circle">'
        '<a:fillToRect l="50000" t="50000" r="50000" b="50000"/></a:path></a:gradFill>'))


def scrim(shape, hex_, alpha):
    """Linear gradient: `alpha` on the left fading to clear on the right."""
    _set_fill_xml(shape, (
        '<a:gradFill rotWithShape="1"><a:gsLst>'
        f'<a:gs pos="0">{_clr(hex_, alpha)}</a:gs>'
        f'<a:gs pos="55000">{_clr(hex_, alpha * 0.55)}</a:gs>'
        f'<a:gs pos="100000">{_clr(hex_, 0.0001)}</a:gs>'
        '</a:gsLst><a:lin ang="0" scaled="0"/></a:gradFill>'))


def no_line(shape):
    shape.line.fill.background()


def line(shape, hex_, width_pt, dash=None, alpha=1.0):
    shape.line.color.rgb = RGBColor.from_string(hex_)
    shape.line.width = Pt(width_pt)
    if dash:
        shape.line.dash_style = dash
    if alpha < 1:
        clr = shape._element.spPr.find(qn("a:ln")).find(".//" + qn("a:srgbClr"))
        etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 100000)))


def set_geometry(pic, prst, adj=None):
    """Give a picture a preset outline (ellipse, roundRect…) instead of a rectangle."""
    geom = pic._element.spPr.find(qn("a:prstGeom"))
    geom.set("prst", prst)
    av = geom.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(geom, qn("a:avLst"))
    for c in list(av):
        av.remove(c)
    if adj is not None:
        gd = etree.SubElement(av, qn("a:gd"))
        gd.set("name", "adj")
        gd.set("fmla", f"val {int(adj)}")


def cover_crop(pic, target_ratio):
    iw, ih = pic.image.size
    ratio = iw / ih
    if ratio > target_ratio:
        c = (1 - target_ratio / ratio) / 2
        pic.crop_left = pic.crop_right = c
    else:
        c = (1 - ratio / target_ratio) / 2
        pic.crop_top = pic.crop_bottom = c


# ---------------------------------------------------------------------------
# Text fitting
# ---------------------------------------------------------------------------

def _wrap_lines(text, chars_per_line):
    lines = 0
    for para in text.split("\n"):
        cur = 0
        lines += 1
        for word in para.split():
            need = len(word) + (1 if cur else 0)
            if cur and cur + need > chars_per_line:
                lines += 1
                cur = len(word)
            else:
                cur += need
    return lines


def fit(text, w, h, max_pt, min_pt=10, wf=0.64, lh=1.18):
    """Largest point size whose estimated wrap fits a w x h inch box.

    `wf` is the average glyph width as a fraction of the point size; Unbounded
    is a wide face, so the defaults are deliberately conservative.
    """
    for pt in range(int(max_pt), int(min_pt) - 1, -1):
        cpl = max(1, int(w / (pt / 72 * wf)))
        if max((len(word) for word in text.split()), default=0) > cpl:
            continue  # never let a single word break mid-way
        if _wrap_lines(text, cpl) * pt / 72 * lh <= h:
            return pt
    return min_pt


class Deck:
    def __init__(self, spec, theme, motion="normal", aspect="16:9"):
        self.spec = spec
        self.t = theme
        self.m = MOTION[motion]
        self.SW, self.SH = ASPECTS[aspect]
        self.sx, self.sy = self.SW / DW, self.SH / DH
        fonts = spec.get("fonts", {})
        self.head_font = fonts.get("heading", "Unbounded")
        self.body_font = fonts.get("body", "Unbounded")
        self.dir = spec.get("_dir", ".")
        self.prs = Presentation()
        self.prs.slide_width = emu(self.SW)
        self.prs.slide_height = emu(self.SH)
        self.blank = self.prs.slide_layouts[6]
        self.layout_count = {}
        self.credits = []

    # -- geometry -----------------------------------------------------------

    def R(self, x, y, w, h, round_=False):
        """Design inches -> slide EMUs. Round shapes keep their aspect ratio."""
        if round_:
            k = min(self.sx, self.sy)
            cx, cy = (x + w / 2) * self.sx, (y + h / 2) * self.sy
            return emu(cx - w * k / 2), emu(cy - h * k / 2), emu(w * k), emu(h * k)
        return emu(x * self.sx), emu(y * self.sy), emu(w * self.sx), emu(h * self.sy)

    def fit(self, text, w, h, *a, **k):
        return fit(text, w * self.sx, h * self.sy, *a, **k)

    # -- shapes -------------------------------------------------------------

    def shape(self, slide, kind, name, x, y, w, h, rot=0, round_=False):
        sh = slide.shapes.add_shape(kind, *self.R(x, y, w, h, round_))
        _strip_style(sh)
        sh.name = name
        sh.rotation = rot
        return sh

    def picture(self, slide, name, ref, x, y, w, h, shape="wide", round_=False):
        try:
            path, credit = images.resolve(ref, self.dir, shape)
        except Exception as e:  # a missing image should not kill the deck
            print(f"warning: image '{ref}' skipped: {e}", file=sys.stderr)
            return None, None
        X, Y, W, H = self.R(x, y, w, h, round_)
        pic = slide.shapes.add_picture(path, X, Y, W, H)
        pic.name = name
        cover_crop(pic, W / H)
        return pic, credit

    def text(self, slide, name, x, y, w, h, paras, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, shape=None):
        """paras: list of dicts {text, size, bold, color, font, spc, alpha, lh, after}."""
        if shape is None:
            shape = slide.shapes.add_textbox(*self.R(x, y, w, h))
            shape.name = name
        tf = shape.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align
            para.line_spacing = p.get("lh", 1.0)
            if p.get("after"):
                para.space_after = Pt(p["after"])
            run = para.add_run()
            run.text = p["text"]
            f = run.font
            f.name = p.get("font", self.body_font)
            f.size = Pt(p["size"])
            f.bold = p.get("bold", False)
            f.color.rgb = RGBColor.from_string(p.get("color", self.t["text"]))
            rpr = run._r.get_or_add_rPr()
            if p.get("spc"):
                rpr.set("spc", str(p["spc"]))
            if p.get("alpha", 1) < 1:
                clr = rpr.find(".//" + qn("a:srgbClr"))
                etree.SubElement(clr, qn("a:alpha")).set("val", str(int(p["alpha"] * 100000)))
        return shape

    def head(self, text, w, h, max_pt, min_pt=18, color=None, lh=1.02):
        size = self.fit(text, w, h, max_pt, min_pt, wf=0.74, lh=lh * 1.12)
        return {"text": text, "size": size, "bold": True, "font": self.head_font,
                "color": color or self.t["text"], "lh": lh}

    def body(self, text, w, h, max_pt=20, min_pt=11, color=None):
        size = self.fit(text, w, h, max_pt, min_pt, wf=0.60, lh=1.45)
        return {"text": text, "size": size, "color": color or self.t["muted"], "lh": 1.25}

    def kicker(self, text, color=None):
        return {"text": text.upper(), "size": 12, "bold": True, "spc": 400,
                "color": color or self.t["accent"], "font": self.head_font}

    def card(self, slide, name, x, y, w, h, paras, fill=None, alpha=1.0, border=True,
             anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, pad=0.32):
        sh = self.shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, name, x, y, w, h)
        sh.adjustments[0] = min(0.12, 0.25 / min(w, h))
        solid(sh, fill or self.t["surface"], alpha)
        if border:
            line(sh, self.t["accent"], 0.75, alpha=0.35)
        else:
            no_line(sh)
        self.text(slide, name, 0, 0, 0, 0, paras, align=align, anchor=anchor, shape=sh)
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = emu(pad * self.sx)
        tf.margin_top = tf.margin_bottom = emu(pad)
        return sh

    def title_block(self, slide, s, a, max_pt=36, w=11.5):
        """Standard top-left slide title, rising in first."""
        h = self.text(slide, "title", 0.9, 0.55, w, 1.15, [self.head(s["title"], w, 1.15, max_pt)],
                      anchor=MSO_ANCHOR.BOTTOM)
        a.add("rise", h, 0, 700)
        return h

    # -- stage --------------------------------------------------------------

    def stage(self, slide, layout, idx, n):
        """Persistent `!!` shapes. Their pose per layout is what Morph animates."""
        t, P = self.t, POSES[layout]
        k = self.layout_count.get(layout, 0)
        spin = self.m["spin"] * idx  # keeps the dashed ring turning across slides
        g = t["glow"]
        glow(self.shape(slide, MSO_SHAPE.OVAL, "!!glow_a", *P["ga"], round_=True), t["accent"], g)
        glow(self.shape(slide, MSO_SHAPE.OVAL, "!!glow_b", *P["gb"], round_=True), t["accent2"], g * 0.8)
        panel = self.shape(slide, MSO_SHAPE.RECTANGLE, "!!panel", *P["panel"])
        solid(panel, t["accent"] if layout in ACCENT_PANEL else t["surface"],
              0.6 if layout in ("cards", "gallery") else 0.9)
        no_line(panel)
        ring = self.shape(slide, MSO_SHAPE.OVAL, "!!ring", *P["ring"], rot=(spin + 37 * k) % 360, round_=True)
        _set_fill_xml(ring, "<a:noFill/>")
        line(ring, t["accent2"], 1.5, dash=MSO_LINE.DASH, alpha=0.7)
        disc = self.shape(slide, MSO_SHAPE.OVAL, "!!disc", *P["disc"], round_=True)
        solid(disc, t["accent2"])
        no_line(disc)
        bar = self.shape(slide, MSO_SHAPE.RECTANGLE, "!!bar", *P["bar"])
        solid(bar, t["accent"])
        no_line(bar)
        prog = self.shape(slide, MSO_SHAPE.RECTANGLE, "!!progress", 0, DH - 0.07,
                          max(0.05, DW * (idx + 1) / n), 0.07)
        solid(prog, t["accent2"], 0.85)
        no_line(prog)
        self.layout_count[layout] = k + 1
        return {"disc": disc, "ring": ring, "bar": bar, "pose": P}

    # -- build --------------------------------------------------------------

    def image_jobs(self):
        """(ref, shape) for every image in the spec; must mirror the layouts' picture() calls."""
        jobs = []
        for s in self.spec["slides"]:
            lay = s.get("layout")
            if lay == "gallery":
                imgs = s.get("images", [])[:4]
                for it in imgs:
                    ref = it if isinstance(it, str) else it.get("image")
                    jobs.append((ref, "tall" if len(imgs) > 2 else "wide"))
            elif lay == "people":
                jobs += [(p["image"], "square") for p in s.get("people", [])[:4] if p.get("image")]
            elif s.get("image"):
                shape = {"image": "wide", "quote": "square"}.get(lay, "tall")
                jobs.append((s["image"], shape))
        return [j for j in jobs if j[0]]

    def build(self):
        images.prefetch(self.image_jobs(), self.dir)
        slides = self.spec["slides"]
        n = len(slides)
        for idx, s in enumerate(slides):
            layout = s.get("layout", "bullets")
            if layout not in LAYOUTS:
                sys.exit(f"slide {idx + 1}: unknown layout '{layout}' (use one of {', '.join(LAYOUTS)})")
            slide = self.prs.slides.add_slide(self.blank)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(self.t["bg"])
            stage = self.stage(slide, layout, idx, n)
            anim = Anim(self.m)
            self.credits = []
            getattr(self, "L_" + layout)(slide, s, anim, stage)
            # ambient loops on the stage: gentle float + a slow sway
            anim.add("float", stage["disc"], 200, 2600)
            anim.add("sway", stage["ring"], 0, 6000)
            notes = "\n\n".join(filter(None, [s.get("notes")] + self.credits))
            if notes:
                slide.notes_slide.notes_text_frame.text = notes
            if idx:
                add_morph(slide, int(self.t["tempo"] * self.m["tempo"] * 1000))
            anim.attach(slide)
        return self.prs

    def _credit(self, credit):
        if credit:
            self.credits.append("Image credit: " + credit)

    # -- layouts ------------------------------------------------------------

    def L_title(self, slide, s, a, st):
        t = self.t
        if s.get("kicker"):
            k = self.text(slide, "kicker", 0.9, 1.55, 8.0, 0.4, [self.kicker(s["kicker"])])
            a.add("fade", k, 0, 600)
        h = self.text(slide, "title", 0.9, 1.95, 8.3, 2.95, [self.head(s["title"], 8.3, 2.95, 66)],
                      anchor=MSO_ANCHOR.BOTTOM)
        a.add("letters", h, 150, 450)
        if s.get("subtitle"):
            sub = self.text(slide, "subtitle", 0.9, 5.4, 7.6, 1.3, [self.body(s["subtitle"], 7.6, 1.3, 20)])
            a.add("rise", sub, 900, 700)
        if s.get("byline"):
            b = self.text(slide, "byline", 0.9, 6.75, 7.6, 0.4,
                          [{"text": s["byline"], "size": 11, "color": t["muted"]}])
            a.add("fade", b, 1300, 600)

    def L_section(self, slide, s, a, st):
        t = self.t
        num = s.get("number", "")
        if num:
            nb = self.text(slide, "number", 0.9, 1.25, 5.0, 5.0,
                           [{"text": num, "size": self.fit(num, 4.2, 3.0, 150, 60, wf=0.8), "bold": True,
                             "font": self.head_font, "color": t["accent"]}],
                           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            a.add("zoom", nb, 0, 700)
        h = self.text(slide, "title", 6.6, 1.3, 6.0, 2.8, [self.head(s["title"], 6.0, 2.8, 50)],
                      anchor=MSO_ANCHOR.BOTTOM)
        a.add("letters", h, 300, 450)
        if s.get("subtitle"):
            sub = self.text(slide, "subtitle", 6.6, 4.6, 5.9, 1.6, [self.body(s["subtitle"], 5.9, 1.6, 18)])
            a.add("rise", sub, 1000, 700)

    def L_agenda(self, slide, s, a, st):
        t = self.t
        items = s.get("items") or [x.get("title", "") for x in self.spec["slides"] if x.get("layout") == "section"]
        items = items[:7]
        h = self.text(slide, "title", 0.9, 1.0, 4.6, 2.2, [self.head(s.get("title", "Agenda"), 4.6, 2.2, 44)],
                      anchor=MSO_ANCHOR.BOTTOM)
        a.add("letters", h, 0, 450)
        row = min(1.15, 5.6 / max(1, len(items)))
        top = 3.75 - row * len(items) / 2
        size = min([self.fit(it, 5.2, min(row - 0.2, 1.0), 24, 12, wf=0.70) for it in items] or [20])
        for i, it in enumerate(items):
            y = top + i * row
            num = self.text(slide, f"num{i}", 6.4, y, 0.9, row, [
                {"text": f"{i + 1:02d}", "size": max(12, int(size * 0.75)), "bold": True,
                 "font": self.head_font, "color": t["accent"] if i % 2 == 0 else t["accent2"]}],
                anchor=MSO_ANCHOR.MIDDLE)
            tx = self.text(slide, f"item{i}", 7.35, y, 5.2, row, [
                {"text": it, "size": size, "bold": True, "font": self.head_font, "color": t["text"]}],
                anchor=MSO_ANCHOR.MIDDLE)
            d = 500 + i * 200
            a.add("fade", num, d, 500)
            a.add("rise", tx, d + 80, 600)

    def L_statement(self, slide, s, a, st):
        t = self.t
        q = self.text(slide, "statement", 1.95, 2.05, 9.6, 3.2,
                      [self.head(s["text"], 9.6, 3.1, 42, 20, lh=1.1)], anchor=MSO_ANCHOR.MIDDLE)
        a.add("rise", q, 0, 900, by_word=True)
        if s.get("attribution"):
            at = self.text(slide, "attribution", 1.95, 5.45, 9.4, 0.5,
                           [{"text": "— " + s["attribution"], "size": 14, "color": t["accent"],
                             "font": self.head_font}])
            a.add("fade", at, 1300, 700)

    def L_question(self, slide, s, a, st):
        t = self.t
        k = self.text(slide, "kicker", 1.5, 1.2, 10.33, 0.4, [self.kicker(s.get("kicker", "Think about it"))],
                      align=PP_ALIGN.CENTER)
        a.add("fade", k, 0, 600)
        q = self.text(slide, "question", 1.5, 1.75, 10.33, 3.3,
                      [self.head(s["text"], 10.3, 3.3, 54, 22, lh=1.08)],
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        a.add("rise", q, 250, 900, by_word=True)
        if s.get("subtitle"):
            sub = self.text(slide, "subtitle", 2.5, 5.55, 8.33, 1.0, [self.body(s["subtitle"], 8.3, 1.0, 18)],
                            align=PP_ALIGN.CENTER)
            a.add("fade", sub, 1500, 700)

    def L_quote(self, slide, s, a, st):
        t = self.t
        has_photo = bool(s.get("image"))
        w = 7.4
        mark = self.text(slide, "mark", 0.75, 0.35, 2.0, 1.6,
                         [{"text": "“", "size": 120, "bold": True, "font": self.head_font, "color": t["accent"]}])
        a.add("zoom", mark, 0, 600)
        q = self.text(slide, "quote", 0.9, 1.65, w, 3.45,
                      [self.head(s["text"], w, 3.45, 38, 18, lh=1.12)], anchor=MSO_ANCHOR.BOTTOM)
        a.add("rise", q, 250, 900, by_word=True)
        who = [{"text": s.get("author", ""), "size": 16, "bold": True, "font": self.head_font,
                "color": t["text"], "after": 4}]
        if s.get("role"):
            who.append({"text": s["role"], "size": 13, "color": t["muted"]})
        au = self.text(slide, "author", 0.9, 5.55, w, 1.1, who)
        a.add("rise", au, 1300, 600)
        if has_photo:
            ring = st["pose"]["ring"]
            x, y, d = ring[0] + 0.35, ring[1] + 0.35, ring[2] - 0.7
            pic, credit = self.picture(slide, "photo", s["image"], x, y, d, d, shape="square", round_=True)
            if pic is not None:
                set_geometry(pic, "ellipse")
                self._credit(credit)
                a.add("zoom", pic, 500, 900)

    def L_bullets(self, slide, s, a, st):
        t = self.t
        h = self.text(slide, "title", 0.9, 0.7, 7.6, 1.2, [self.head(s["title"], 7.6, 1.2, 36)],
                      anchor=MSO_ANCHOR.BOTTOM)
        a.add("rise", h, 0, 700)
        pts = s.get("points", [])[:6]
        top = 2.45
        row = min(1.25, 4.4 / max(1, len(pts)))
        size = min(self.fit(p, 6.9, row - 0.15, 20, 11, wf=0.60, lh=1.4) for p in pts) if pts else 18
        for i, p in enumerate(pts):
            y = top + i * row
            dot = self.shape(slide, MSO_SHAPE.OVAL, f"dot{i}", 0.92, y + size / 72 * 0.42, 0.14, 0.14, round_=True)
            solid(dot, t["accent2"] if i % 2 else t["accent"])
            no_line(dot)
            tx = self.text(slide, f"point{i}", 1.35, y, 6.9, row - 0.1,
                           [{"text": p, "size": size, "color": t["text"], "lh": 1.2}])
            a.add("zoom", dot, 500 + i * 220, 400)
            a.add("rise", tx, 550 + i * 220, 600)
        if s.get("image"):
            pic, credit = self.picture(slide, "image", s["image"], 8.9, 0, DW - 8.9, DH, shape="tall")
            if pic is not None:
                self._credit(credit)
                a.add("fade", pic, 200, 1000)
        elif s.get("aside"):
            asd = self.text(slide, "aside", 9.5, 2.6, 3.4, 2.4,
                            [self.head(s["aside"], 3.4, 2.4, 30, 14, color=t["text"])],
                            anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
            a.add("zoom", asd, 400 + len(pts) * 220, 700)

    def L_stat(self, slide, s, a, st):
        t = self.t
        value, suffix = str(s.get("value", "")), str(s.get("suffix", ""))
        size = self.fit(value + suffix, 6.8, 3.0, 170, 48, wf=0.82, lh=1.0)
        v = self.text(slide, "value", 0.9, 1.0, 7.2, 3.9, [], anchor=MSO_ANCHOR.BOTTOM)
        para = v.text_frame.paragraphs[0]
        para.line_spacing = 0.9
        for txt, col, sz in ((value, t["text"], size), (suffix, t["accent"], int(size * 0.7))):
            if not txt:
                continue
            r = para.add_run()
            r.text = txt
            r.font.name, r.font.size, r.font.bold = self.head_font, Pt(sz), True
            r.font.color.rgb = RGBColor.from_string(col)
        a.add("zoom", v, 0, 800)
        lbl = self.text(slide, "label", 0.9, 5.3, 6.8, 0.75, [self.head(s.get("label", ""), 6.8, 0.75, 24, 14)])
        a.add("rise", lbl, 600, 600)
        if s.get("note"):
            nt = self.text(slide, "note", 0.9, 6.05, 6.8, 1.0, [self.body(s["note"], 6.8, 1.0, 15)])
            a.add("fade", nt, 1000, 700)
        if s.get("context"):
            c = self.text(slide, "context", 8.9, 2.75, 3.6, 2.0,
                          [self.body(s["context"], 3.6, 2.0, 18, 11, color=t["text"])],
                          anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
            a.add("fade", c, 1200, 800)

    def L_chart(self, slide, s, a, st):
        t = self.t
        self.title_block(slide, s, a, w=8.2)
        kind = s.get("type", "column")
        types = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
                 "line": XL_CHART_TYPE.LINE_MARKERS, "donut": XL_CHART_TYPE.DOUGHNUT,
                 "pie": XL_CHART_TYPE.PIE, "area": XL_CHART_TYPE.AREA,
                 "stacked": XL_CHART_TYPE.COLUMN_STACKED}
        if kind not in types:
            sys.exit(f"chart type '{kind}' not supported (use one of {', '.join(types)})")
        cd = CategoryChartData()
        cd.categories = s["categories"]
        series = s["series"]
        for se in series:
            cd.add_series(se.get("name", ""), se["values"], number_format=s.get("number_format", "General"))
        gf = slide.shapes.add_chart(types[kind], *self.R(0.75, 2.0, 8.4, 4.85), cd)
        gf.name = "chart"
        ch = gf.chart
        ch.has_title = False
        palette = [t["accent"], t["accent2"], t["text"], t["muted"]]
        ch.font.name = self.body_font
        ch.font.size = Pt(12)
        ch.font.color.rgb = RGBColor.from_string(t["muted"])
        round_kind = kind in ("donut", "pie")
        ch.has_legend = len(series) > 1 or round_kind
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
            ch.legend.font.color.rgb = RGBColor.from_string(t["text"])
        plot = ch.plots[0]
        if round_kind:
            pts = plot.series[0].points
            for i in range(len(s["categories"])):
                pts[i].format.fill.solid()
                pts[i].format.fill.fore_color.rgb = RGBColor.from_string(
                    [t["accent"], t["accent2"], t["muted"], t["text"], t["surface"]][i % 5])
                pts[i].format.line.color.rgb = RGBColor.from_string(t["bg"])
                pts[i].format.line.width = Pt(2)
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.show_value = True
            dl.font.size, dl.font.bold = Pt(13), True
            dl.font.color.rgb = RGBColor.from_string(t["bg"])
            dl.number_format = s.get("number_format", "0")
            dl.number_format_is_linked = False
            if kind == "donut":
                hole = plot._element.find(qn("c:holeSize"))
                if hole is None:
                    hole = etree.SubElement(plot._element, qn("c:holeSize"))
                hole.set("val", "62")
        else:
            for i, se in enumerate(plot.series):
                col = RGBColor.from_string(palette[i % len(palette)])
                if kind == "line":
                    se.format.line.color.rgb = col
                    se.format.line.width = Pt(3.5)
                    se.smooth = True
                    se.marker.format.fill.solid()
                    se.marker.format.fill.fore_color.rgb = col
                    se.marker.format.line.color.rgb = col
                else:
                    se.format.fill.solid()
                    se.format.fill.fore_color.rgb = col
            if kind in ("column", "bar", "stacked"):
                plot.gap_width = 70
                if kind != "stacked":
                    plot.overlap = -10
            if len(series) == 1 and kind in ("column", "bar"):
                plot.has_data_labels = True
                plot.data_labels.font.size = Pt(12)
                plot.data_labels.font.bold = True
                plot.data_labels.font.color.rgb = RGBColor.from_string(t["text"])
                plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
            va, ca = ch.value_axis, ch.category_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = RGBColor.from_string(t["muted"])
            va.major_gridlines.format.line.width = Pt(0.5)
            va.format.line.fill.background()
            va.tick_labels.font.color.rgb = RGBColor.from_string(t["muted"])
            va.tick_labels.number_format = s.get("axis_format", "General")
            va.tick_labels.number_format_is_linked = False
            ca.format.line.color.rgb = RGBColor.from_string(t["muted"])
            ca.tick_labels.font.color.rgb = RGBColor.from_string(t["text"])
            ca.tick_labels.font.size = Pt(12)
        anim = {"column": ("wipe", "up"), "stacked": ("wipe", "up"), "bar": ("wipe", "right"),
                "line": ("wipe", "right"), "area": ("wipe", "right"),
                "donut": ("wheel", None), "pie": ("wheel", None)}[kind]
        a.add(anim[0], gf, 500, 1400, dir=anim[1])
        if s.get("takeaway"):
            tk = self.text(slide, "takeaway", 9.75, 2.4, 2.4, 2.4,
                           [self.head(s["takeaway"], 2.4, 2.4, 40, 16, color=t["text"])],
                           anchor=MSO_ANCHOR.BOTTOM)
            a.add("zoom", tk, 1600, 700)
        if s.get("caption"):
            cp = self.text(slide, "caption", 9.75, 4.95, 2.4, 1.6, [self.body(s["caption"], 2.4, 1.6, 14, 10)])
            a.add("fade", cp, 1900, 700)
        if s.get("source"):
            src = self.text(slide, "source", 0.9, 6.95, 8.0, 0.35,
                            [{"text": "Source: " + s["source"], "size": 9, "color": t["muted"]}])
            a.add("fade", src, 2000, 500)

    def L_cards(self, slide, s, a, st):
        t = self.t
        self.title_block(slide, s, a)
        cards = s.get("cards", [])[:4]
        gap, x0, x1, y, hgt = 0.3, 0.9, DW - 0.9, 2.35, 4.15
        w = (x1 - x0 - gap * (len(cards) - 1)) / max(1, len(cards))
        inner = w - 0.64
        tsize = min(self.fit(c.get("title", ""), inner, 1.1, 22, 13, wf=0.70) for c in cards) if cards else 20
        bsize = min(self.fit(c.get("text", ""), inner, 2.4, 18, 10, wf=0.60, lh=1.45) for c in cards) if cards else 14
        for i, c in enumerate(cards):
            paras = [{"text": f"{i + 1:02d}", "size": 14, "bold": True, "font": self.head_font,
                      "color": t["accent"], "after": 14},
                     {"text": c.get("title", ""), "size": tsize, "bold": True, "font": self.head_font,
                      "color": t["text"], "after": 10, "lh": 1.05},
                     {"text": c.get("text", ""), "size": bsize, "color": t["muted"], "lh": 1.25}]
            cd = self.card(slide, f"card{i}", x0 + i * (w + gap), y, w, hgt, paras)
            a.add("rise", cd, 500 + i * 250, 700)

    def L_people(self, slide, s, a, st):
        t = self.t
        self.title_block(slide, s, a)
        ppl = s.get("people", [])[:4]
        n = max(1, len(ppl))
        col = (DW - 1.8) / n
        d = min(2.3, col - 0.8)
        nsize = min([self.fit(p.get("name", ""), col - 0.4, 0.6, 20, 12, wf=0.72) for p in ppl] or [18])
        for i, p in enumerate(ppl):
            cx = 0.9 + col * i + col / 2
            x, y = cx - d / 2, 2.25
            pic = None
            if p.get("image"):
                pic, credit = self.picture(slide, f"photo{i}", p["image"], x, y, d, d, shape="square", round_=True)
                if pic is not None:
                    set_geometry(pic, "ellipse")
                    self._credit(credit)
            if pic is None:  # initials disc
                initials = "".join(w[0] for w in p.get("name", "?").split()[:2]).upper()
                pic = self.shape(slide, MSO_SHAPE.OVAL, f"photo{i}", x, y, d, d, round_=True)
                solid(pic, t["surface"])
                line(pic, t["accent"] if i % 2 == 0 else t["accent2"], 2)
                size = int(d * min(self.sx, self.sy) * 72 * 0.55 / (0.8 * max(1, len(initials))))
                self.text(slide, "", 0, 0, 0, 0, [{"text": initials, "size": min(size, 60), "bold": True,
                                                    "font": self.head_font, "color": t["text"]}],
                          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=pic)
                pic.text_frame.word_wrap = False
            paras = [{"text": p.get("name", ""), "size": nsize, "bold": True, "font": self.head_font,
                      "color": t["text"], "after": 4},
                     {"text": p.get("role", ""), "size": 13, "color": t["accent"], "after": 8}]
            if p.get("text"):
                paras.append({"text": p["text"], "size": self.fit(p["text"], col - 0.4, 1.3, 14, 10, wf=0.6, lh=1.45),
                              "color": t["muted"], "lh": 1.25})
            tx = self.text(slide, f"person{i}", cx - (col - 0.4) / 2, y + d + 0.3, col - 0.4, 2.4, paras,
                           align=PP_ALIGN.CENTER)
            dl = 500 + i * 250
            a.add("zoom", pic, dl, 700)
            a.add("rise", tx, dl + 200, 600)

    def L_timeline(self, slide, s, a, st):
        t = self.t
        self.title_block(slide, s, a)
        steps = s.get("steps", [])[:6]
        n = max(1, len(steps))
        span = 11.6 / n
        lsize = min(self.fit(x.get("label", ""), span - 0.3, 0.9, 30, 12, wf=0.80) for x in steps) if steps else 24
        bsize = min(self.fit(x.get("text", ""), span - 0.3, 2.0, 15, 10, wf=0.60, lh=1.45) for x in steps) if steps else 13
        for i, x in enumerate(steps):
            cx = 0.9 + span * i + 0.1
            node = self.shape(slide, MSO_SHAPE.OVAL, f"node{i}", cx, 4.115, 0.3, 0.3, round_=True)
            solid(node, t["bg"])
            line(node, t["accent"], 2.25)
            lb = self.text(slide, f"label{i}", cx, 2.85, span - 0.3, 1.0,
                           [{"text": x.get("label", ""), "size": lsize, "bold": True,
                             "font": self.head_font, "color": t["accent"] if i % 2 == 0 else t["accent2"]}],
                           anchor=MSO_ANCHOR.BOTTOM)
            tx = self.text(slide, f"step{i}", cx, 4.75, span - 0.3, 2.1,
                           [{"text": x.get("text", ""), "size": bsize, "color": t["text"], "lh": 1.25}])
            d = 500 + i * 380
            a.add("zoom", node, d, 400)
            a.add("rise", lb, d + 120, 600)
            a.add("fade", tx, d + 250, 700)

    def L_process(self, slide, s, a, st):
        """Steps listed on the left, numbered nodes orbiting the stage ring on the right."""
        t = self.t
        h = self.text(slide, "title", 0.9, 0.6, 6.3, 1.25, [self.head(s["title"], 6.3, 1.25, 34)],
                      anchor=MSO_ANCHOR.BOTTOM)
        a.add("rise", h, 0, 700)
        steps = s.get("steps", [])[:5]
        n = max(1, len(steps))
        rx, ry, rw, _ = st["pose"]["ring"]
        k = min(self.sx, self.sy)
        ccx, ccy, r = (rx + rw / 2) * self.sx, (ry + rw / 2) * self.sy, rw / 2 * k
        top, bottom = 2.35, 6.85
        row = (bottom - top) / n
        lsize = min(self.fit(x.get("label", ""), 5.9, 0.45, 20, 11, wf=0.72) for x in steps) if steps else 16
        bsize = min(self.fit(x.get("text", ""), 6.1, row - 0.5, 16, 10, wf=0.6, lh=1.4) for x in steps) if steps else 12
        for i, x in enumerate(steps):
            ang = -math.pi / 2 + 2 * math.pi * i / n
            nd = 0.62
            nx, ny = (ccx + r * math.cos(ang)) / self.sx, (ccy + r * math.sin(ang)) / self.sy
            col = t["accent"] if i % 2 == 0 else t["accent2"]
            node = self.shape(slide, MSO_SHAPE.OVAL, f"node{i}", nx - nd / 2, ny - nd / 2, nd, nd, round_=True)
            solid(node, col)
            no_line(node)
            self.text(slide, "", 0, 0, 0, 0, [{"text": str(i + 1), "size": 16, "bold": True,
                                               "font": self.head_font, "color": t["bg"]}],
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shape=node)
            y = top + i * row
            tx = self.text(slide, f"step{i}", 0.9, y, 6.2, row - 0.1, [
                {"text": f"{i + 1:02d}  " + x.get("label", ""), "size": lsize, "bold": True,
                 "font": self.head_font, "color": col, "after": 4},
                {"text": x.get("text", ""), "size": bsize, "color": t["text"], "lh": 1.2}])
            d = 500 + i * 330
            a.add("zoom", node, d, 450)
            a.add("rise", tx, d + 100, 600)
        if s.get("center"):
            c = self.text(slide, "center", (ccx - r * 0.62) / self.sx, (ccy - r * 0.4) / self.sy,
                          r * 1.24 / self.sx, r * 0.8 / self.sy,
                          [self.head(s["center"], r * 1.24 / self.sx, r * 0.8 / self.sy, 24, 12)],
                          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            a.add("fade", c, 500 + n * 330, 700)

    def L_compare(self, slide, s, a, st):
        t = self.t
        self.title_block(slide, s, a)
        for side, x, col, d in (("left", 0.9, t["accent"], 500), ("right", 7.15, t["accent2"], 800)):
            c = s.get(side, {})
            pts = c.get("points", [])
            bsize = min([self.fit(p, 4.7, 3.4 / max(1, len(pts)), 17, 10, wf=0.60, lh=1.4) for p in pts] or [16])
            paras = [{"text": c.get("title", "").upper(), "size": 14, "bold": True, "spc": 300,
                      "font": self.head_font, "color": col, "after": 16}]
            paras += [{"text": "—  " + p, "size": bsize, "color": t["text"], "after": 10, "lh": 1.2} for p in pts]
            cd = self.card(slide, side, x, 2.15, 5.3, 4.5, paras, border=False, pad=0.4)
            a.add("rise", cd, d, 700)

    def L_split(self, slide, s, a, st):
        t = self.t
        if s.get("kicker"):
            k = self.text(slide, "kicker", 0.9, 1.1, 5.6, 0.4, [self.kicker(s["kicker"])])
            a.add("fade", k, 0, 500)
        h = self.text(slide, "title", 0.9, 1.5, 5.6, 1.7, [self.head(s["title"], 5.6, 1.7, 36)],
                      anchor=MSO_ANCHOR.BOTTOM)
        a.add("rise", h, 100, 700)
        if s.get("text"):
            b = self.text(slide, "body", 0.9, 3.65, 5.6, 3.2, [self.body(s["text"], 5.6, 3.2, 18, 11)])
            a.add("rise", b, 600, 700)
        if s.get("image"):
            pic, credit = self.picture(slide, "image", s["image"], 7.0, 0, DW - 7.0, DH, shape="tall")
            if pic is not None:
                self._credit(credit)
                a.add("fade", pic, 300, 1000)
                a.add("kenburns", pic, 300, 9000)
        elif s.get("aside"):
            asd = self.text(slide, "aside", 8.0, 2.4, 3.6, 2.7,
                            [self.head(s["aside"], 3.6, 2.7, 34, 14)],
                            anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
            a.add("zoom", asd, 900, 700)

    def L_image(self, slide, s, a, st):
        """Full-bleed photo with a theme-coloured scrim and text over the lower left."""
        t = self.t
        pic, credit = self.picture(slide, "image", s["image"], 0, 0, DW, DH, shape="wide")
        if pic is not None:
            self._credit(credit)
            a.add("fade", pic, 0, 900)
            a.add("kenburns", pic, 0, 10000)
        sc = self.shape(slide, MSO_SHAPE.RECTANGLE, "scrim", 0, 0, 9.5, DH)
        scrim(sc, t["bg"], 0.92)
        no_line(sc)
        bar = self.shape(slide, MSO_SHAPE.RECTANGLE, "accent", 0.9, 4.6, 1.0, 0.09)
        solid(bar, t["accent"])
        no_line(bar)
        a.add("wipe", bar, 500, 500, dir="right")
        if s.get("kicker"):
            k = self.text(slide, "kicker", 0.9, 2.25, 6.5, 0.4, [self.kicker(s["kicker"])])
            a.add("fade", k, 300, 500)
        h = self.text(slide, "title", 0.9, 2.65, 6.8, 1.8, [self.head(s["title"], 6.8, 1.8, 44)],
                      anchor=MSO_ANCHOR.BOTTOM)
        a.add("letters", h, 400, 450)
        if s.get("text"):
            b = self.text(slide, "body", 0.9, 4.95, 6.0, 1.9,
                          [self.body(s["text"], 6.0, 1.9, 18, 11, color=t["text"])])
            a.add("rise", b, 1200, 700)

    def L_gallery(self, slide, s, a, st):
        t = self.t
        self.title_block(slide, s, a)
        items = s.get("images", [])[:4]
        n = max(1, len(items))
        gap, x0, y, hgt = 0.3, 0.9, 2.2, 3.9
        w = (DW - 1.8 - gap * (n - 1)) / n
        for i, it in enumerate(items):
            it = {"image": it} if isinstance(it, str) else it
            x = x0 + i * (w + gap)
            pic, credit = self.picture(slide, f"img{i}", it["image"], x, y, w, hgt,
                                       shape="tall" if n > 2 else "wide")
            if pic is None:
                continue
            set_geometry(pic, "roundRect", adj=min(16667, 0.25 / min(w, hgt) * 50000))
            self._credit(credit)
            d = 500 + i * 220
            a.add("zoom", pic, d, 800)
            if it.get("caption"):
                cp = self.text(slide, f"cap{i}", x, y + hgt + 0.18, w, 0.6,
                               [{"text": it["caption"], "size": self.fit(it["caption"], w, 0.6, 14, 10, wf=0.62),
                                 "color": t["muted"]}])
                a.add("fade", cp, d + 300, 600)

    def L_closing(self, slide, s, a, st):
        t = self.t
        h = self.text(slide, "title", 1.4, 1.6, 10.53, 2.6, [self.head(s["title"], 10.5, 2.6, 60)],
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
        a.add("letters", h, 100, 450)
        if s.get("subtitle"):
            sub = self.text(slide, "subtitle", 2.2, 4.75, 8.93, 1.3, [self.body(s["subtitle"], 8.9, 1.3, 20)],
                            align=PP_ALIGN.CENTER)
            a.add("rise", sub, 900, 700)


# ---------------------------------------------------------------------------
# Transition + animation XML
# ---------------------------------------------------------------------------

def add_morph(slide, dur_ms):
    xml = (
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
        f'<p:transition {NS_DECL} xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" '
        f'spd="slow" p14:dur="{dur_ms}"><p159:morph option="byObject"/></p:transition>'
        '</mc:Choice><mc:Fallback>'
        f'<p:transition {NS_DECL} spd="slow"><p:fade/></p:transition>'
        '</mc:Fallback></mc:AlternateContent>')
    el = etree.fromstring(xml)
    sld = slide._element
    anchor = sld.find(qn("p:timing"))
    if anchor is None:
        anchor = sld.find(qn("p:extLst"))
    if anchor is not None:
        anchor.addprevious(el)
    else:
        sld.append(el)


class Anim:
    """Collects effects; all start automatically once the slide (and its Morph) lands.

    Each effect is "With Previous" plus its own delay, which is how PowerPoint
    stores a staggered After-Previous chain, so it stays editable in the
    Animation Pane. The motion profile scales delays, durations and distances.
    """

    PRESETS = {  # kind: (presetClass, presetID, presetSubtype)
        "fade": ("entr", 10, 0), "rise": ("entr", 42, 0), "letters": ("entr", 42, 0),
        "zoom": ("entr", 53, 16), "wipe": ("entr", 22, 8), "wheel": ("entr", 21, 1),
        "float": ("path", 42, 0), "sway": ("emph", 8, 0), "kenburns": ("emph", 6, 0),
    }
    ENTRANCE = ("fade", "rise", "letters", "zoom", "wipe", "wheel")
    WIPE_SUB = {"right": 8, "left": 2, "up": 4, "down": 1}

    def __init__(self, motion):
        self.m = motion
        self.effects = []

    def add(self, kind, shape, delay, dur=600, **opts):
        if shape is None:
            return
        if kind == "letters" and not self.m["letters"]:
            kind = "rise"
        if kind in self.ENTRANCE:
            delay, dur = delay * self.m["stagger"], dur * self.m["dur"]
        self.effects.append((kind, shape, int(delay), int(dur), opts))

    def attach(self, slide):
        if not self.effects:
            return
        ids = itertools.count(5)
        pars, blds = [], []
        for kind, shape, delay, dur, o in self.effects:
            pars.append(self._effect(kind, shape.shape_id, delay, dur, o, ids))
            if kind not in self.ENTRANCE:
                continue
            if shape.__class__.__name__ == "GraphicFrame":
                blds.append(f'<p:bldGraphic spid="{shape.shape_id}" grpId="0"><p:bldAsOne/></p:bldGraphic>')
            elif getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                anim_bg = "" if kind == "letters" or o.get("by_word") else ' animBg="1"'
                blds.append(f'<p:bldP spid="{shape.shape_id}" grpId="0"{anim_bg}/>')
        bld = f'<p:bldLst>{"".join(dict.fromkeys(blds))}</p:bldLst>' if blds else ""
        xml = (
            f'<p:timing {NS_DECL}><p:tnLst><p:par>'
            '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
            '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
            '<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
            '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
            '<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            + "".join(pars) +
            '</p:childTnLst></p:cTn></p:par>'
            '</p:childTnLst></p:cTn></p:par>'
            '</p:childTnLst></p:cTn>'
            '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
            '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
            '</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>' + bld + '</p:timing>')
        old = slide._element.find(qn("p:timing"))
        if old is not None:
            slide._element.remove(old)
        el = etree.fromstring(xml)
        ext = slide._element.find(qn("p:extLst"))
        if ext is not None:
            ext.addprevious(el)
        else:
            slide._element.append(el)

    @staticmethod
    def _tgt(spid):
        return f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'

    def _set_visible(self, spid, ids):
        return (f'<p:set><p:cBhvr><p:cTn id="{next(ids)}" dur="1" fill="hold"><p:stCondLst>'
                f'<p:cond delay="0"/></p:stCondLst></p:cTn>{self._tgt(spid)}'
                '<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
                '<p:to><p:strVal val="visible"/></p:to></p:set>')

    def _filter(self, spid, flt, dur, ids):
        return (f'<p:animEffect transition="in" filter="{flt}"><p:cBhvr><p:cTn id="{next(ids)}" dur="{dur}"/>'
                f'{self._tgt(spid)}</p:cBhvr></p:animEffect>')

    def _anim(self, spid, attr, frm, to, dur, ids):
        return (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base">'
                f'<p:cTn id="{next(ids)}" dur="{dur}" decel="100000" fill="hold"/>{self._tgt(spid)}'
                f'<p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst>'
                f'<p:tav tm="0"><p:val><p:strVal val="{frm}"/></p:val></p:tav>'
                f'<p:tav tm="100000"><p:val><p:strVal val="{to}"/></p:val></p:tav></p:tavLst></p:anim>')

    def _effect(self, kind, spid, delay, dur, o, ids):
        cls, pid, sub = self.PRESETS[kind]
        eid = next(ids)
        extra, iterate = "", ""
        loop = ' accel="50000" decel="50000" repeatCount="indefinite" autoRev="1"'
        if kind == "letters" or o.get("by_word"):
            unit = "lt" if kind == "letters" else "wd"
            iterate = f'<p:iterate type="{unit}"><p:tmPct val="{6000 if unit == "lt" else 12000}"/></p:iterate>'
        if kind == "float":
            extra = loop
            body = (f'<p:animMotion origin="layout" path="M 0 0 L 0 {-self.m["float"]} E" pathEditMode="relative" '
                    f'ptsTypes=""><p:cBhvr><p:cTn id="{next(ids)}" dur="{dur}" fill="hold"/>{self._tgt(spid)}'
                    '<p:attrNameLst><p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst>'
                    '</p:cBhvr></p:animMotion>')
        elif kind == "sway":
            extra = loop
            body = (f'<p:animRot by="{int(self.m["sway"] * 60000)}"><p:cBhvr><p:cTn id="{next(ids)}" dur="{dur}" '
                    f'fill="hold"/>{self._tgt(spid)}<p:attrNameLst><p:attrName>r</p:attrName></p:attrNameLst>'
                    '</p:cBhvr></p:animRot>')
        elif kind == "kenburns":
            extra = ' decel="100000"'
            sc = self.m["kenburns"]
            body = (f'<p:animScale><p:cBhvr><p:cTn id="{next(ids)}" dur="{dur}" fill="hold"/>{self._tgt(spid)}'
                    f'</p:cBhvr><p:by x="{sc}" y="{sc}"/></p:animScale>')
        else:
            body = self._set_visible(spid, ids)
            if kind == "wipe":
                d = o.get("dir") or "right"
                sub = self.WIPE_SUB[d]
                body += self._filter(spid, f"wipe({d})", dur, ids)
            elif kind == "wheel":
                body += self._filter(spid, "wheel(1)", dur, ids)
            else:
                body += self._filter(spid, "fade", dur, ids)
            if kind in ("rise", "letters"):
                body += self._anim(spid, "ppt_y", f"#ppt_y+{self.m['dy']}", "#ppt_y", dur, ids)
            if kind == "zoom":
                z = self.m["zoom"]
                body += self._anim(spid, "ppt_w", f"#ppt_w*{z}", "#ppt_w", dur, ids)
                body += self._anim(spid, "ppt_h", f"#ppt_h*{z}", "#ppt_h", dur, ids)
        return (f'<p:par><p:cTn id="{eid}" presetID="{pid}" presetClass="{cls}" presetSubtype="{sub}"{extra} '
                f'fill="hold" grpId="0" nodeType="withEffect"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>'
                f'{iterate}<p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>')


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec", nargs="?")
    ap.add_argument("out", nargs="?")
    ap.add_argument("--theme", help="override the theme named in the spec")
    ap.add_argument("--motion", choices=list(MOTION), help="override options.motion")
    ap.add_argument("--aspect", choices=list(ASPECTS), help="override options.aspect")
    ap.add_argument("--list-themes", action="store_true")
    args = ap.parse_args()
    themes = load_themes()
    if args.list_themes:
        for k, v in themes.items():
            print(f"{k:10s} {'dark ' if v['dark'] else 'light'}  {v['mood']}")
        return
    if not (args.spec and args.out):
        ap.error("spec and out are required")
    with open(args.spec) as f:
        spec = json.load(f)
    spec["_dir"] = os.path.dirname(os.path.abspath(args.spec))
    opts = spec.get("options", {})
    name = args.theme or spec.get("theme", "midnight")
    if name not in themes:
        sys.exit(f"unknown theme '{name}'. Available: {', '.join(themes)}")
    motion = args.motion or opts.get("motion", "normal")
    aspect = args.aspect or opts.get("aspect", "16:9")
    if motion not in MOTION:
        sys.exit(f"unknown motion '{motion}'. Use one of: {', '.join(MOTION)}")
    if aspect not in ASPECTS:
        sys.exit(f"unsupported aspect '{aspect}'. Use one of: {', '.join(ASPECTS)}")
    theme = dict(themes[name], **spec.get("theme_overrides", {}))
    prs = Deck(spec, theme, motion, aspect).build()
    prs.core_properties.title = spec.get("title", "")
    prs.save(args.out)
    want = opts.get("slides")
    got = len(spec["slides"])
    note = f" (options.slides asked for {want})" if want and want != got else ""
    print(f"wrote {args.out}: {got} slides{note}, theme '{name}', motion '{motion}', aspect {aspect}")


if __name__ == "__main__":
    main()
